# -*- coding: utf-8 -*-
"""
PythonAnywhere WSGI 入口
在 PythonAnywhere → Web → WSGI 配置文件中指向此文件
"""
import os, json, threading, io, urllib.parse
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 数据文件路径（PythonAnywhere 文件系统）
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, "data.json")
LOCK = threading.Lock()


def _find_site_dir():
    """在多个可能位置查找 site 目录"""
    candidates = [
        os.path.join(BASE_DIR, "site"),
        os.path.join(os.path.dirname(BASE_DIR), "site"),
        BASE_DIR,
    ]
    for c in candidates:
        if os.path.isfile(os.path.join(c, "index.html")):
            return c
    return os.path.join(BASE_DIR, "site")


SITE_DIR = _find_site_dir()


def _load_seed():
    """从 seed.js 解析种子数据"""
    candidates = [
        os.path.join(SITE_DIR, "js", "seed.js"),
        os.path.join(BASE_DIR, "site", "js", "seed.js"),
        os.path.join(BASE_DIR, "js", "seed.js"),
    ]
    for seed_path in candidates:
        if os.path.exists(seed_path):
            try:
                with open(seed_path, "r", encoding="utf-8") as f:
                    content = f.read()
                start = content.index("{")
                end = content.rindex("}") + 1
                return json.loads(content[start:end])
            except Exception:
                pass
    return {"settings": {"companyName": "中咨养护检测", "currentMonth": "2026-09", "dailyRate": 300},
            "projects": [], "personnel": [], "achievements": [],
            "targets": {}, "monthlyReports": {}, "signatures": {}}


SEED = _load_seed()


def load_data():
    with LOCK:
        if os.path.exists(DATA_FILE):
            try:
                with open(DATA_FILE, "r", encoding="utf-8") as f:
                    return json.load(f)
            except Exception:
                pass
        # 初始化种子数据
        with open(DATA_FILE, "w", encoding="utf-8") as f:
            json.dump(SEED, f, ensure_ascii=False)
        return json.loads(json.dumps(SEED))


def save_data(data):
    with LOCK:
        with open(DATA_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False)


def generate_attendance_excel(data, month, pid="all"):
    """生成考勤表 Excel"""
    settings = data.get("settings", {})
    projects = data.get("projects", [])
    personnel_all = data.get("personnel", [])

    year, mon = int(month[:4]), int(month[5:7])
    start_month = mon - 2
    start_year = year
    if start_month <= 0:
        start_month += 12
        start_year -= 1
    start_date = date(start_year, start_month, 26)
    end_month = mon - 1
    end_year = year
    if end_month <= 0:
        end_month += 12
        end_year -= 1
    end_date = date(end_year, end_month, 25)

    dates = []
    d = start_date
    while d <= end_date:
        dates.append(d)
        d += timedelta(days=1)
    n_days = len(dates)

    if pid and pid != "all":
        personnel = [pe for pe in personnel_all if pe.get("projectId") == pid]
    else:
        personnel = list(personnel_all)

    proj_name = "全部项目"
    if pid and pid != "all":
        for p in projects:
            if p.get("id") == pid:
                proj_name = p.get("name", "")
                break

    att_month = "%d年%d月" % (end_year, end_month)

    wb = Workbook()
    ws = wb.active
    ws.title = att_month

    F_TITLE = Font(name="宋体", size=16)
    F_TEXT = Font(name="宋体", size=10.5)
    F_UNIT = Font(name="宋体", size=11)
    F_HEAD = Font(name="宋体", size=10)
    F_DATE = Font(name="仿宋_GB2312", size=9)
    F_DATA = Font(name="宋体", size=10)
    F_NAME = Font(name="宋体", size=9)
    F_DAILY = Font(name="仿宋_GB2312", size=9)
    F_SUM = Font(name="宋体", size=10)
    F_SIGN = Font(name="隶书", size=11)

    A_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    A_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    A_LEFT_TOP = Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="000000")
    BORDER_ALL = Border(top=thin, bottom=thin, left=thin, right=thin)
    BORDER_TOP_LEFT = Border(top=thin, left=thin, right=thin)

    total_cols = 6 + n_days * 2

    ws.cell(row=1, column=1, value="研发人员考勤明细表")
    ws.cell(row=1, column=1).font = F_TITLE
    ws.cell(row=1, column=1).alignment = A_CENTER
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws.row_dimensions[1].height = 30

    ws.row_dimensions[2].height = 8

    time_col_start = 23
    time_col_end = time_col_start + 7
    ws.cell(row=3, column=time_col_start, value="时间： %d  年 %d月" % (end_year, end_month))
    ws.cell(row=3, column=time_col_start).font = F_TEXT
    ws.cell(row=3, column=time_col_start).alignment = A_CENTER
    ws.merge_cells(start_row=3, start_column=time_col_start, end_row=3, end_column=time_col_end)
    ws.row_dimensions[3].height = 20

    ws.cell(row=4, column=1, value="项目名称：%s" % proj_name)
    ws.cell(row=4, column=1).font = F_TEXT
    ws.cell(row=4, column=1).alignment = A_LEFT
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=3)

    unit_col = 65
    ws.cell(row=4, column=unit_col, value="单位：天")
    ws.cell(row=4, column=unit_col).font = F_UNIT
    ws.cell(row=4, column=unit_col).alignment = A_LEFT_TOP
    ws.row_dimensions[4].height = 20

    headers = ["序号", "姓名", "职责分工", "研发课题名称", "研发人员研发工时（天）", "研发人员总工时（天）"]
    for c, hdr in enumerate(headers):
        col = c + 1
        cell = ws.cell(row=5, column=col, value=hdr)
        cell.font = F_HEAD
        cell.alignment = A_CENTER
        cell.border = BORDER_TOP_LEFT
        ws.merge_cells(start_row=5, start_column=col, end_row=7, end_column=col)

    weekday_names = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
    for di, d in enumerate(dates):
        col_start = 7 + di * 2
        label = "%d日" % d.day
        cell = ws.cell(row=5, column=col_start, value=label)
        cell.font = F_DATE
        cell.alignment = A_CENTER
        cell.border = Border(top=thin, left=thin)
        ws.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_start + 1)

    for di, d in enumerate(dates):
        col_start = 7 + di * 2
        wd = weekday_names[d.weekday()]
        cell = ws.cell(row=6, column=col_start, value=wd)
        cell.font = F_DATE
        cell.alignment = A_CENTER
        cell.border = Border(left=thin)
        ws.merge_cells(start_row=6, start_column=col_start, end_row=6, end_column=col_start + 1)

    for di in range(n_days):
        col_start = 7 + di * 2
        cell_rd = ws.cell(row=7, column=col_start, value="研发人员工时")
        cell_rd.font = F_DATE
        cell_rd.alignment = A_CENTER
        cell_rd.border = BORDER_ALL
        cell_nrd = ws.cell(row=7, column=col_start + 1, value="非研发人员工时")
        cell_nrd.font = F_DATE
        cell_nrd.alignment = A_CENTER
        cell_nrd.border = BORDER_ALL

    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 18
    ws.row_dimensions[7].height = 25

    data_start_row = 8
    sum_rd_total = 0
    sum_tot_total = 0
    daily_rd_sums = [0] * n_days
    daily_nrd_sums = [0] * n_days

    for pi, pe in enumerate(personnel):
        row = data_start_row + pi
        ws.cell(row=row, column=1, value=pi + 1).font = F_DATA
        ws.cell(row=row, column=1).alignment = A_CENTER
        ws.cell(row=row, column=1).border = BORDER_ALL

        ws.cell(row=row, column=2, value=pe.get("name", "")).font = F_NAME
        ws.cell(row=row, column=2).alignment = A_CENTER
        ws.cell(row=row, column=2).border = BORDER_ALL

        ws.cell(row=row, column=3, value=pe.get("responsibility", "")).font = F_DATA
        ws.cell(row=row, column=3).alignment = A_CENTER
        ws.cell(row=row, column=3).border = BORDER_ALL

        pe_proj_name = ""
        pe_pid = pe.get("projectId")
        if pe_pid:
            for p in projects:
                if p.get("id") == pe_pid:
                    pe_proj_name = p.get("name", "")
                    break
        ws.cell(row=row, column=4, value=pe_proj_name).font = F_DATA
        ws.cell(row=row, column=4).alignment = A_CENTER
        ws.cell(row=row, column=4).border = BORDER_ALL

        att = pe.get("attendance", {}).get(month, {})
        days_data = att.get("days", {})

        rd_sum = 0
        tot_sum = 0
        for di, d in enumerate(dates):
            pos_key = str(di + 1)
            day_data = days_data.get(pos_key, {})
            if not day_data:
                day_data = days_data.get(di + 1, {})
            rd_val = float(day_data.get("rd", 0) or 0)
            nrd_val = float(day_data.get("nonRd", 0) or 0)
            rd_sum += rd_val
            tot_sum += rd_val + nrd_val
            daily_rd_sums[di] += rd_val
            daily_nrd_sums[di] += nrd_val

            col_start = 7 + di * 2
            cell_rd = ws.cell(row=row, column=col_start, value=round(rd_val, 1))
            cell_rd.font = F_DAILY
            cell_rd.alignment = A_CENTER
            cell_rd.border = BORDER_ALL

            cell_nrd = ws.cell(row=row, column=col_start + 1, value=round(nrd_val, 1))
            cell_nrd.font = F_DAILY
            cell_nrd.alignment = A_CENTER
            cell_nrd.border = BORDER_ALL

        ws.cell(row=row, column=5, value=round(rd_sum, 1)).font = F_DATA
        ws.cell(row=row, column=5).alignment = A_CENTER
        ws.cell(row=row, column=5).border = BORDER_ALL

        ws.cell(row=row, column=6, value=round(tot_sum, 1)).font = F_DATA
        ws.cell(row=row, column=6).alignment = A_CENTER
        ws.cell(row=row, column=6).border = BORDER_ALL

        sum_rd_total += rd_sum
        sum_tot_total += tot_sum
        ws.row_dimensions[row].height = 18

    sum_row = data_start_row + len(personnel)
    ws.cell(row=sum_row, column=1, value="合计").font = F_SUM
    ws.cell(row=sum_row, column=1).alignment = A_CENTER
    ws.cell(row=sum_row, column=1).border = BORDER_ALL
    ws.merge_cells(start_row=sum_row, start_column=1, end_row=sum_row, end_column=3)
    for c in range(2, 4):
        ws.cell(row=sum_row, column=c).border = BORDER_ALL

    ws.cell(row=sum_row, column=4, value="").border = BORDER_ALL
    ws.cell(row=sum_row, column=5, value=round(sum_rd_total, 1)).font = F_SUM
    ws.cell(row=sum_row, column=5).alignment = A_CENTER
    ws.cell(row=sum_row, column=5).border = BORDER_ALL
    ws.cell(row=sum_row, column=6, value=round(sum_tot_total, 1)).font = F_SUM
    ws.cell(row=sum_row, column=6).alignment = A_CENTER
    ws.cell(row=sum_row, column=6).border = BORDER_ALL

    for di in range(n_days):
        col_start = 7 + di * 2
        cell_rd = ws.cell(row=sum_row, column=col_start, value=round(daily_rd_sums[di], 1))
        cell_rd.font = F_SUM
        cell_rd.alignment = A_CENTER
        cell_rd.border = BORDER_ALL
        cell_nrd = ws.cell(row=sum_row, column=col_start + 1, value=round(daily_nrd_sums[di], 1))
        cell_nrd.font = F_SUM
        cell_nrd.alignment = A_CENTER
        cell_nrd.border = BORDER_ALL
    ws.row_dimensions[sum_row].height = 20

    empty_row = sum_row + 1
    ws.row_dimensions[empty_row].height = 8

    sign_row = sum_row + 2
    sign_col1 = 21
    sig_key = month + (("_" + pid) if pid and pid != "all" else "")
    signatures = settings.get("signatures", {})
    sig = signatures.get(sig_key, {})
    ws.cell(row=sign_row, column=sign_col1, value="课题负责人：  %s" % (sig.get("leaderName", "")))
    ws.cell(row=sign_row, column=sign_col1).font = F_SIGN
    ws.cell(row=sign_row, column=sign_col1).alignment = A_LEFT_TOP

    sign_col2 = 63
    ws.cell(row=sign_row, column=sign_col2, value="制表人：  %s" % (sig.get("makerName", "")))
    ws.cell(row=sign_row, column=sign_col2).font = F_SIGN
    ws.cell(row=sign_row, column=sign_col2).alignment = A_LEFT_TOP
    ws.row_dimensions[sign_row].height = 25

    ws.column_dimensions["A"].width = 6.5
    ws.column_dimensions["B"].width = 6.5
    ws.column_dimensions["C"].width = 6.5
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 8
    for di in range(n_days):
        col1 = get_column_letter(7 + di * 2)
        col2 = get_column_letter(7 + di * 2 + 1)
        ws.column_dimensions[col1].width = 5
        ws.column_dimensions[col2].width = 5

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def serve_static(environ, start_response):
    """提供静态文件"""
    path_info = environ.get("PATH_INFO", "/")
    # 映射到 site 目录
    if path_info == "/" or path_info == "":
        path_info = "/index.html"

    # 安全：防止路径遍历
    rel = path_info.lstrip("/").replace("\\", "/")
    if ".." in rel.split("/"):
        start_response("403 Forbidden", [("Content-Type", "text/plain")])
        return [b"403 Forbidden"]

    file_path = os.path.join(SITE_DIR, rel)

    if not os.path.exists(file_path) or not os.path.isfile(file_path):
        start_response("404 Not Found", [("Content-Type", "text/plain; charset=utf-8")])
        return [b"404 Not Found: " + rel.encode("utf-8")]

    ext = os.path.splitext(file_path)[1].lower()
    content_types = {
        ".html": "text/html; charset=utf-8",
        ".css": "text/css; charset=utf-8",
        ".js": "application/javascript; charset=utf-8",
        ".json": "application/json; charset=utf-8",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".gif": "image/gif",
        ".svg": "image/svg+xml",
        ".ico": "image/x-icon",
    }
    ct = content_types.get(ext, "application/octet-stream")

    try:
        with open(file_path, "rb") as f:
            content = f.read()
        start_response("200 OK", [
            ("Content-Type", ct),
            ("Content-Length", str(len(content))),
            ("Access-Control-Allow-Origin", "*"),
        ])
        return [content]
    except Exception as e:
        start_response("500 Internal Server Error", [("Content-Type", "text/plain")])
        return [str(e).encode("utf-8")]


def application(environ, start_response):
    """WSGI 入口"""
    path_info = environ.get("PATH_INFO", "/")
    method = environ.get("REQUEST_METHOD", "GET")

    # CORS 预检
    if method == "OPTIONS":
        start_response("200 OK", [
            ("Access-Control-Allow-Origin", "*"),
            ("Access-Control-Allow-Methods", "GET, POST, OPTIONS"),
            ("Access-Control-Allow-Headers", "Content-Type"),
        ])
        return [b""]

    # API 路由
    if path_info == "/api/data" and method == "GET":
        try:
            data = load_data()
            body = json.dumps(data, ensure_ascii=False).encode("utf-8")
            start_response("200 OK", [
                ("Content-Type", "application/json; charset=utf-8"),
                ("Content-Length", str(len(body))),
                ("Access-Control-Allow-Origin", "*"),
            ])
            return [body]
        except Exception as e:
            start_response("500 Internal Server Error", [("Content-Type", "application/json")])
            return [json.dumps({"error": str(e)}).encode("utf-8")]

    elif path_info == "/api/data" and method == "POST":
        try:
            content_length = int(environ.get("CONTENT_LENGTH", 0))
            body_bytes = environ["wsgi.input"].read(content_length) if content_length > 0 else b""
            data = json.loads(body_bytes.decode("utf-8"))
            save_data(data)
            resp = json.dumps({"ok": True}).encode("utf-8")
            start_response("200 OK", [
                ("Content-Type", "application/json; charset=utf-8"),
                ("Content-Length", str(len(resp))),
                ("Access-Control-Allow-Origin", "*"),
            ])
            return [resp]
        except Exception as e:
            start_response("500 Internal Server Error", [("Content-Type", "application/json")])
            return [json.dumps({"error": str(e)}).encode("utf-8")]

    elif path_info == "/api/reset" and method == "GET":
        try:
            data = json.loads(json.dumps(SEED))
            save_data(data)
            resp = json.dumps({"ok": True}).encode("utf-8")
            start_response("200 OK", [
                ("Content-Type", "application/json; charset=utf-8"),
                ("Access-Control-Allow-Origin", "*"),
            ])
            return [resp]
        except Exception as e:
            start_response("500 Internal Server Error", [("Content-Type", "application/json")])
            return [json.dumps({"error": str(e)}).encode("utf-8")]

    elif path_info == "/api/diag" and method == "GET":
        # 诊断接口：返回路径和文件检查信息
        diag = {
            "ok": True,
            "base_dir": BASE_DIR,
            "site_dir": SITE_DIR,
            "data_file": DATA_FILE,
            "data_file_exists": os.path.exists(DATA_FILE),
            "files": {},
            "seed_projects": len(SEED.get("projects", [])),
            "seed_people": len(SEED.get("people", SEED.get("personnel", []))),
        }
        for rel in ["index.html", "css/styles.css", "js/app.js", "js/seed.js", "vendor/xlsx.full.min.js"]:
            fp = os.path.join(SITE_DIR, rel)
            diag["files"][rel] = {"exists": os.path.exists(fp), "size": os.path.getsize(fp) if os.path.exists(fp) else 0}
        try:
            d = load_data()
            diag["loaded_projects"] = len(d.get("projects", []))
            diag["loaded_personnel"] = len(d.get("personnel", d.get("people", [])))
        except Exception as e:
            diag["load_error"] = str(e)
        body = json.dumps(diag, ensure_ascii=False, indent=2).encode("utf-8")
        start_response("200 OK", [
            ("Content-Type", "application/json; charset=utf-8"),
            ("Content-Length", str(len(body))),
            ("Access-Control-Allow-Origin", "*"),
        ])
        return [body]

    elif path_info.startswith("/api/export/attendance"):
        try:
            query = urllib.parse.parse_qs(environ.get("QUERY_STRING", ""))
            month = query.get("month", [""])[0]
            pid = query.get("pid", ["all"])[0]
            if not month:
                start_response("400 Bad Request", [("Content-Type", "application/json")])
                return [json.dumps({"error": "缺少 month 参数"}).encode("utf-8")]
            data = load_data()
            body = generate_attendance_excel(data, month, pid)
            filename = "研发人员考勤表-%s.xlsx" % month
            start_response("200 OK", [
                ("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                ("Content-Disposition", "attachment; filename*=UTF-8''" + urllib.parse.quote(filename)),
                ("Content-Length", str(len(body))),
                ("Access-Control-Allow-Origin", "*"),
            ])
            return [body]
        except Exception as e:
            start_response("500 Internal Server Error", [("Content-Type", "application/json")])
            return [json.dumps({"error": str(e)}).encode("utf-8")]

    # 静态文件
    return serve_static(environ, start_response)
